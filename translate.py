#!/usr/bin/env python3
"""Add Estonian (et-EE) translations to MyHealth@EU cross-border value sets.

Reads every value set workbook (.xlsx) in input/, looks up each SNOMED CT
concept in the Estonian SNOMED CT extension found in sources/, writes the
Estonian preferred term into column E and saves the result in output/.

Usage:
    python translate.py [--input input] [--sources sources] [--output output]
"""
import argparse
import collections
import csv
import fnmatch
import io
import re
import sys
import warnings
import zipfile
from pathlib import Path

import openpyxl

SNOMED_OID = "2.16.840.1.113883.6.96"
LANGUAGE = "et"
LANGUAGE_REFSET = "71000181105"  # Estonian language reference set
PREFERRED = "900000000000548007"
NOT_AVAILABLE = "N/A"

HEADER_LABEL = "Code System ID"  # column A label of the concept table header row
COL_SYSTEM, COL_CODE, COL_TERM, COL_TRANSLATION = 1, 3, 4, 5
COLUMN_WIDTHS = {"B": 20, "D": 50, "E": 50}

DESCRIPTION_PATTERN = f"*Snapshot/Terminology/sct2_Description_Snapshot-{LANGUAGE}_*.txt"
LANGUAGE_PATTERN = f"*Snapshot/Refset/Language/der2_cRefset_LanguageSnapshot-{LANGUAGE}_*.txt"

csv.field_size_limit(sys.maxsize)
warnings.filterwarnings("ignore", message="Workbook contains no default style")


class Release:
    """An RF2 release given as a .zip file or an extracted folder."""

    def __init__(self, path):
        self.path = path
        if path.is_dir():
            self.names = [p.relative_to(path).as_posix() for p in path.rglob("*") if p.is_file()]
        else:
            with zipfile.ZipFile(path) as zf:
                self.names = zf.namelist()

    def find(self, pattern):
        hits = [n for n in self.names if fnmatch.fnmatch(n, pattern)]
        if len(hits) != 1:
            raise SystemExit(f"{self.path.name}: expected one file matching {pattern}, found {len(hits)}")
        return hits[0]

    def rows(self, pattern):
        name = self.find(pattern)
        if self.path.is_dir():
            fh = open(self.path / name, encoding="utf-8", newline="")
        else:
            fh = io.TextIOWrapper(zipfile.ZipFile(self.path).open(name), encoding="utf-8", newline="")
        with fh:
            yield from csv.DictReader(fh, delimiter="\t", quoting=csv.QUOTE_NONE)


def find_snomed_release(sources):
    candidates = [p for p in sorted(sources.iterdir()) if p.name.startswith("SnomedCT_")
                  and (p.is_dir() or p.suffix.lower() == ".zip")]
    if not candidates:
        raise SystemExit(f"No SNOMED CT release (SnomedCT_*.zip or folder) found in {sources}/")
    # Release names end in the release date (…_YYYYMMDDT…); take the newest.
    candidates.sort(key=lambda p: re.findall(r"\d{8}T", p.name)[-1:] or [""])
    if len(candidates) > 1:
        print(f"Several SNOMED CT releases in {sources}/, using the newest: {candidates[-1].name}")
    return candidates[-1]


def load_preferred_terms(release):
    """Map conceptId -> Estonian preferred term (active descriptions, active refset members)."""
    descriptions = {r["id"]: r for r in release.rows(DESCRIPTION_PATTERN) if r["active"] == "1"}
    terms = collections.defaultdict(set)
    for r in release.rows(LANGUAGE_PATTERN):
        if (r["active"] == "1" and r["refsetId"] == LANGUAGE_REFSET
                and r["acceptabilityId"] == PREFERRED and r["referencedComponentId"] in descriptions):
            d = descriptions[r["referencedComponentId"]]
            terms[d["conceptId"]].add(d["term"])
    for concept, found in terms.items():
        if len(found) > 1:
            print(f"  warning: concept {concept} has {len(found)} preferred terms, using the first "
                  f"alphabetically: {sorted(found)}")
    return {concept: sorted(found)[0] for concept, found in terms.items()}


def header_row(ws):
    for cell in ws["A"]:
        if isinstance(cell.value, str) and cell.value.strip() == HEADER_LABEL:
            return cell.row
    return None


def translate_sheet(ws, translations):
    """Fill column E for SNOMED CT rows; return counts per outcome."""
    counts = collections.Counter()
    start = header_row(ws)
    if start is None:
        counts["skipped sheet (no concept table)"] = 1
        return counts
    for row in range(start + 1, ws.max_row + 1):
        system = ws.cell(row, COL_SYSTEM).value
        code = ws.cell(row, COL_CODE).value
        if system is None and code is None:
            continue
        if str(system).strip() != SNOMED_OID:
            counts["other code system (left unchanged)"] += 1
            continue
        term = translations.get(str(code).strip(), NOT_AVAILABLE)
        ws.cell(row, COL_TRANSLATION).value = term
        counts["N/A" if term == NOT_AVAILABLE else "translated"] += 1
    for column, width in COLUMN_WIDTHS.items():
        ws.column_dimensions[column].width = width
    return counts


def output_name(sheet_names):
    name = "MVC_" + "_".join(sheet_names) + f"_{LANGUAGE}-EE.xlsx"
    if len(name) > 150:
        name = f"MVC_{sheet_names[0]}_and_{len(sheet_names) - 1}_more_{LANGUAGE}-EE.xlsx"
    return name


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    here = Path(__file__).resolve().parent
    parser.add_argument("--input", type=Path, default=here / "input")
    parser.add_argument("--sources", type=Path, default=here / "sources")
    parser.add_argument("--output", type=Path, default=here / "output")
    args = parser.parse_args()

    workbooks = sorted(p for p in args.input.glob("*.xlsx") if not p.name.startswith("~$"))
    if not workbooks:
        raise SystemExit(f"No .xlsx value sets found in {args.input}/")

    release = Release(find_snomed_release(args.sources))
    print(f"SNOMED CT release: {release.path.name}")
    translations = load_preferred_terms(release)
    print(f"Estonian preferred terms loaded for {len(translations)} concepts\n")

    args.output.mkdir(parents=True, exist_ok=True)
    for path in workbooks:
        wb = openpyxl.load_workbook(path)
        value_sets = wb.worksheets[1:]  # first sheet is the catalogue info, left unchanged
        print(path.name)
        for ws in value_sets:
            counts = translate_sheet(ws, translations)
            print(f"  {ws.title}: " + ", ".join(f"{k} {v}" for k, v in counts.items()))
        target = args.output / output_name([ws.title for ws in value_sets])
        wb.save(target)
        print(f"  -> {target}\n")


if __name__ == "__main__":
    main()
